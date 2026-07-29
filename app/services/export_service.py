"""Exportación a Excel (sección 16 de la especificación).

Requisitos cubiertos: encabezados, formato monetario, encabezados
congelados, autofiltro, ancho de columna ajustado, fecha/hora de
generación, filtros utilizados, hoja de resumen y de detalle, nombres de
archivo seguros, y saneamiento de valores que Excel podría interpretar
como fórmulas (sección 21).
"""
from __future__ import annotations

import dataclasses
import datetime as dt
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Movement
from app.repositories.movement_repository import MovementFilter
from app.schemas.report_filter import ReportFilter
from app.services.report_service import ConsolidatedReport, ReportRow
from app.utils.dates import format_date_for_display
from app.utils.money import cents_to_decimal
from app.utils.text import sanitize_for_spreadsheet

MONEY_FORMAT = "#,##0.00"

HEADER_FILL = PatternFill(start_color="1E5FBF", end_color="1E5FBF", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
DAY_FILL = PatternFill(start_color="DCE6F5", end_color="DCE6F5", fill_type="solid")
BOLD = Font(bold=True)

_UNSAFE_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*]')


def build_safe_filename(
    prefix: str, date_from: dt.date | None = None, date_to: dt.date | None = None
) -> str:
    """Genera un nombre de archivo seguro. Ej: reporte_sucursales_2026-05-01_2026-05-31.xlsx"""
    safe_prefix = _UNSAFE_FILENAME_CHARS.sub("_", prefix).strip().replace(" ", "_")
    parts = [safe_prefix]
    if date_from:
        parts.append(date_from.isoformat())
    if date_to:
        parts.append(date_to.isoformat())
    return "_".join(parts) + ".xlsx"


def describe_filters(filters) -> str:
    """Representación legible de los filtros usados, para incluir en el
    archivo exportado (sección 16: "Incluir los filtros utilizados")."""
    parts = []
    for f in dataclasses.fields(filters):
        value = getattr(filters, f.name)
        if value in (None, "", []):
            continue
        if isinstance(value, dt.date):
            value = format_date_for_display(value)
        parts.append(f"{f.name}={value}")
    return "; ".join(parts) if parts else "Sin filtros"


def _write_header_info(
    ws: Worksheet, title: str, generated_at: dt.datetime, filters_description: str
) -> int:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {generated_at.strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A3"] = f"Filtros: {filters_description}"
    return 5  # fila donde empieza la tabla de datos


def _autosize_columns(ws: Worksheet, columns_count: int, min_width: int = 10, max_width: int = 50) -> None:
    for col_index in range(1, columns_count + 1):
        letter = get_column_letter(col_index)
        max_len = min_width
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, max_width)


def _money(cents: int | None) -> float | None:
    if not cents:
        return None
    return float(cents_to_decimal(cents))


# ---------------------------------------------------------------------------
# Movimientos filtrados
# ---------------------------------------------------------------------------
def export_movements_to_excel(
    movements: list[Movement],
    filters: MovementFilter,
    output_path: Path,
    generated_at: dt.datetime | None = None,
) -> Path:
    generated_at = generated_at or dt.datetime.now()
    workbook = Workbook()

    summary_ws = workbook.active
    summary_ws.title = "Resumen"
    _write_header_info(summary_ws, "Resumen de movimientos", generated_at, describe_filters(filters))

    total_charge = sum(m.charge_cents for m in movements)
    total_payment = sum(m.payment_cents for m in movements)
    summary_ws["A6"] = "Total de movimientos"
    summary_ws["B6"] = len(movements)
    summary_ws["A7"] = "Cargo total"
    summary_ws["B7"] = float(cents_to_decimal(total_charge))
    summary_ws["B7"].number_format = MONEY_FORMAT
    summary_ws["A8"] = "Abono total"
    summary_ws["B8"] = float(cents_to_decimal(total_payment))
    summary_ws["B8"].number_format = MONEY_FORMAT
    _autosize_columns(summary_ws, 2)

    detail_ws = workbook.create_sheet("Detalle")
    headers = [
        "Fecha", "Banco", "Cuenta", "Sucursal", "Categoría",
        "Concepto", "Referencia", "Cargo", "Abono", "Saldo",
        "Clasificación", "Archivo",
    ]
    detail_ws.append(headers)
    for cell in detail_ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for movement in movements:
        detail_ws.append(
            [
                format_date_for_display(movement.movement_date),
                movement.bank_account.bank.name if movement.bank_account else "",
                movement.bank_account.alias if movement.bank_account else "",
                movement.branch.name if movement.branch else "",
                movement.category.name if movement.category else "",
                sanitize_for_spreadsheet(movement.description_original),
                sanitize_for_spreadsheet(movement.reference_original or ""),
                _money(movement.charge_cents),
                _money(movement.payment_cents),
                _money(movement.balance_cents),
                movement.classification_status,
                movement.imported_file.original_name if movement.imported_file else "",
            ]
        )

    for row in detail_ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            cell.number_format = MONEY_FORMAT

    detail_ws.freeze_panes = "A2"
    detail_ws.auto_filter.ref = detail_ws.dimensions
    _autosize_columns(detail_ws, len(headers))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Reportes simples (totales por día/banco/cuenta/sucursal/categoría, etc.)
# ---------------------------------------------------------------------------
def export_simple_report_to_excel(
    title: str,
    rows: list[ReportRow],
    filters: ReportFilter,
    output_path: Path,
    generated_at: dt.datetime | None = None,
) -> Path:
    generated_at = generated_at or dt.datetime.now()
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Reporte"

    start_row = _write_header_info(ws, title, generated_at, describe_filters(filters))

    headers = ["Concepto", "Cargo", "Abono", "Neto"]
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    current_row = start_row + 1
    total_charge = total_payment = 0
    for row in rows:
        ws.cell(row=current_row, column=1, value=sanitize_for_spreadsheet(row.label))
        for col_index, cents in ((2, row.charge_cents), (3, row.payment_cents), (4, row.net_cents)):
            cell = ws.cell(row=current_row, column=col_index, value=float(cents_to_decimal(cents)))
            cell.number_format = MONEY_FORMAT
        total_charge += row.charge_cents
        total_payment += row.payment_cents
        current_row += 1

    last_data_row = current_row - 1
    ws.cell(row=current_row, column=1, value="TOTAL GENERAL").font = BOLD
    for col_index, cents in (
        (2, total_charge),
        (3, total_payment),
        (4, total_payment - total_charge),
    ):
        cell = ws.cell(row=current_row, column=col_index, value=float(cents_to_decimal(cents)))
        cell.number_format = MONEY_FORMAT
        cell.font = BOLD

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate
    if last_data_row >= start_row:
        ws.auto_filter.ref = f"A{start_row}:D{last_data_row}"
    _autosize_columns(ws, len(headers))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# Tabla consolidada (día / sucursal / cuenta)
# ---------------------------------------------------------------------------
def export_consolidated_report_to_excel(
    report: ConsolidatedReport,
    value_type_label: str,
    filters: ReportFilter,
    output_path: Path,
    generated_at: dt.datetime | None = None,
) -> Path:
    generated_at = generated_at or dt.datetime.now()
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Consolidado"

    start_row = _write_header_info(
        ws, "Tabla consolidada por fecha, sucursal y cuenta", generated_at, describe_filters(filters)
    )
    ws["A4"] = f"Valor mostrado: {value_type_label}"

    headers = ["Día / Sucursal", *report.account_aliases, "Total"]
    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    current_row = start_row + 1
    for day_group in report.day_groups:
        day_totals = day_group.values_by_account
        ws.cell(row=current_row, column=1, value=format_date_for_display(day_group.day))
        for col_index, alias in enumerate(report.account_aliases, start=2):
            ws.cell(row=current_row, column=col_index, value=_money(day_totals.get(alias, 0)))
        ws.cell(row=current_row, column=len(headers), value=_money(day_group.total_cents))
        for cell in ws[current_row]:
            cell.font = BOLD
            cell.fill = DAY_FILL
            if cell.column > 1:
                cell.number_format = MONEY_FORMAT
        current_row += 1

        for branch_row in day_group.branch_rows:
            label_cell = ws.cell(row=current_row, column=1, value="    " + branch_row.branch_label)
            label_cell.alignment = Alignment(indent=1)
            for col_index, alias in enumerate(report.account_aliases, start=2):
                cell = ws.cell(
                    row=current_row,
                    column=col_index,
                    value=_money(branch_row.values_by_account.get(alias)),
                )
                cell.number_format = MONEY_FORMAT
            total_cell = ws.cell(row=current_row, column=len(headers), value=_money(branch_row.total_cents))
            total_cell.number_format = MONEY_FORMAT
            current_row += 1

    total_row = current_row
    ws.cell(row=total_row, column=1, value="TOTAL GENERAL").font = BOLD
    grand_totals = report.grand_totals_by_account
    for col_index, alias in enumerate(report.account_aliases, start=2):
        cell = ws.cell(row=total_row, column=col_index, value=_money(grand_totals.get(alias, 0)))
        cell.number_format = MONEY_FORMAT
        cell.font = BOLD
    grand_total_cell = ws.cell(row=total_row, column=len(headers), value=_money(report.grand_total_cents))
    grand_total_cell.number_format = MONEY_FORMAT
    grand_total_cell.font = BOLD

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{total_row}"
    _autosize_columns(ws, len(headers))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
