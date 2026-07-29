import datetime as dt

import pytest
from openpyxl import load_workbook

from app.database.seed import seed_catalogs
from app.models import BankAccount, Branch, ImportedFile, Movement, MovementCategory
from app.repositories.movement_repository import MovementFilter
from app.schemas.report_filter import ReportFilter
from app.services import export_service
from app.services.report_service import ConsolidatedReport, ConsolidatedBranchRow, ConsolidatedDayGroup, ReportRow


def test_build_safe_filename_strips_unsafe_characters():
    name = export_service.build_safe_filename(
        "reporte: sucursales?", dt.date(2026, 5, 1), dt.date(2026, 5, 31)
    )
    assert name == "reporte__sucursales__2026-05-01_2026-05-31.xlsx"
    assert not any(ch in name for ch in '<>:"/\\|?*')


def test_build_safe_filename_without_dates():
    name = export_service.build_safe_filename("movimientos")
    assert name == "movimientos.xlsx"


@pytest.fixture()
def bbva_with_movement(db_session):
    seed_catalogs(db_session)
    bbva = db_session.query(BankAccount).filter_by(alias="BBVA").one()
    branch = db_session.query(Branch).filter_by(branch_number="01").one()
    category = db_session.query(MovementCategory).filter_by(code="BRANCH").one()

    imported_file = ImportedFile(
        bank_account_id=bbva.id,
        original_name="extracto.txt",
        stored_name="extracto.txt",
        original_path="C:/tmp/extracto.txt",
        file_type="TXT",
        file_hash="hash-export-test",
        file_size=10,
    )
    db_session.add(imported_file)
    db_session.flush()

    movement = Movement(
        bank_account_id=bbva.id,
        branch_id=branch.id,
        category_id=category.id,
        imported_file_id=imported_file.id,
        movement_date=dt.date(2026, 7, 1),
        description_original="=CMD|'/c calc'!A1",  # intento de inyección de fórmula
        normalized_text="CMD C CALC A1",
        reference_original="+1234",
        charge_cents=0,
        payment_cents=100000,
        balance_cents=500000,
        movement_hash="hash-export-movement",
        occurrence_number=1,
        classification_status="CLASSIFIED",
        classification_method="BRANCH_IDENTIFIER",
        source_row_number=1,
    )
    db_session.add(movement)
    db_session.flush()
    db_session.refresh(movement)
    # Forzar carga de relaciones mientras la sesión sigue abierta.
    _ = movement.bank_account.bank.name
    _ = movement.branch.name
    _ = movement.category.name
    _ = movement.imported_file.original_name
    return movement


def test_export_movements_to_excel_sanitizes_formula_injection(tmp_path, bbva_with_movement):
    output_path = tmp_path / "movimientos.xlsx"
    export_service.export_movements_to_excel(
        [bbva_with_movement], MovementFilter(), output_path, generated_at=dt.datetime(2026, 7, 29, 10, 0)
    )

    assert output_path.exists()
    workbook = load_workbook(output_path)
    assert "Resumen" in workbook.sheetnames
    assert "Detalle" in workbook.sheetnames

    detail = workbook["Detalle"]
    header_row = [cell.value for cell in detail[1]]
    assert header_row[0] == "Fecha"

    data_row = [cell.value for cell in detail[2]]
    concept_value = data_row[5]
    assert concept_value.startswith("'"), "El valor con '=' debe quedar escapado"

    reference_value = data_row[6]
    assert reference_value.startswith("'"), "El valor con '+' debe quedar escapado"

    summary = workbook["Resumen"]
    assert summary["A1"].value == "Resumen de movimientos"
    assert summary["B6"].value == 1


def test_export_simple_report_to_excel(tmp_path):
    rows = [
        ReportRow(label="BBVA", charge_cents=1000, payment_cents=5000),
        ReportRow(label="MIFEL", charge_cents=0, payment_cents=2000),
    ]
    output_path = tmp_path / "reporte.xlsx"
    export_service.export_simple_report_to_excel(
        "Totales por banco", rows, ReportFilter(), output_path, generated_at=dt.datetime(2026, 7, 29, 10, 0)
    )

    workbook = load_workbook(output_path)
    ws = workbook["Reporte"]
    assert ws["A1"].value == "Totales por banco"
    # La tabla empieza en la fila 5 (después del bloque de encabezado).
    assert ws["A5"].value == "Concepto"
    assert ws["A6"].value == "BBVA"
    assert ws["C6"].value == 50.0
    assert ws["A8"].value == "TOTAL GENERAL"
    assert ws["C8"].value == 70.0


def test_export_consolidated_report_to_excel(tmp_path):
    report = ConsolidatedReport(
        account_aliases=["BBVA", "MIFEL"],
        day_groups=[
            ConsolidatedDayGroup(
                day=dt.date(2026, 7, 1),
                branch_rows=[
                    ConsolidatedBranchRow(
                        branch_label="01. ALFREDO", values_by_account={"BBVA": 100000}
                    ),
                    ConsolidatedBranchRow(
                        branch_label="02. COLON", values_by_account={"MIFEL": 50000}
                    ),
                ],
            )
        ],
    )
    output_path = tmp_path / "consolidado.xlsx"
    export_service.export_consolidated_report_to_excel(
        report, "Abono", ReportFilter(), output_path, generated_at=dt.datetime(2026, 7, 29, 10, 0)
    )

    workbook = load_workbook(output_path)
    ws = workbook["Consolidado"]
    assert ws["A1"].value == "Tabla consolidada por fecha, sucursal y cuenta"
    assert ws["A4"].value == "Valor mostrado: Abono"
    assert ws["A5"].value == "Día / Sucursal"
    assert ws["A6"].value == "01/07/2026"
    assert ws["D6"].value == 1500.0  # total del día (BBVA 1000 + MIFEL 500)
